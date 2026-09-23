from django.http import HttpResponse
from rest_framework import viewsets, permissions

from apps.equipements.models import Equipement
from .models import JournalActivite
from .serializers import JournalActiviteSerializer
from apps.utilisateurs.models import Role

from django.db.models import Q
from rest_framework.pagination import PageNumberPagination

from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from apps.reservations.models import Reservation
from apps.laboratoires.models import Laboratoire

from datetime import datetime
from apps.utilisateurs.models import Role

import openpyxl
from openpyxl.styles import Font, PatternFill

ROLE_LABELS = {
    Role.ADMIN: 'Administrateur', Role.TECHNICIEN: 'Technicien',
    Role.CHERCHEUR: 'Enseignant-chercheur', Role.ETUDIANT: 'Étudiant',
}
class JournalPagination(PageNumberPagination):
    page_size = 9
    page_size_query_param = 'page_size'
    max_page_size = 1000

class EstAdmin(permissions.BasePermission):
    def has_permission(self, request, view):
        return request.user.is_authenticated and request.user.role == Role.ADMIN


class JournalActiviteViewSet(viewsets.ReadOnlyModelViewSet):
    """
    ReadOnlyModelViewSet plutôt que ModelViewSet : un journal ne se crée
    et ne se modifie jamais via l'API, uniquement via journaliser().
    """
    queryset = JournalActivite.objects.all()
    serializer_class = JournalActiviteSerializer
    permission_classes = [EstAdmin]
    pagination_class = JournalPagination

    def get_queryset(self):
        qs = super().get_queryset()
        auteur_id = self.request.query_params.get('auteur')
        action = self.request.query_params.get('action')
        entite = self.request.query_params.get('entite')
        search = self.request.query_params.get('search')
        date_debut = self.request.query_params.get('date_debut')
        date_fin = self.request.query_params.get('date_fin')

        if auteur_id:
            qs = qs.filter(auteur_id=auteur_id)
        if action:
            qs = qs.filter(action__icontains=action)
        if entite:
            qs = qs.filter(entite_type__model=entite)
        if search:
            qs = qs.filter(Q(action__icontains=search) | Q(description__icontains=search))
        if date_debut:
            qs = qs.filter(date_heure__date__gte=date_debut)
        if date_fin:
            qs = qs.filter(date_heure__date__lte=date_fin)
        return qs
    
    
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def recherche_globale(request):
    terme = request.query_params.get('q', '').strip()
    if len(terme) < 2:
        return Response({'equipements': [], 'laboratoires': []})

    equipements = Equipement.objects.filter(nom__icontains=terme)[:5]
    laboratoires = Laboratoire.objects.filter(nom__icontains=terme)[:5]

    return Response({
        'equipements': [{'id': e.id, 'nom': e.nom, 'laboratoire_nom': e.laboratoire.nom} for e in equipements],
        'laboratoires': [{'id': l.id, 'nom': l.nom, 'localisation': l.localisation} for l in laboratoires],
    })
    
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def mon_activite(request):
    entrees = JournalActivite.objects.filter(auteur=request.user).order_by('-date_heure')[:10]
    return Response(JournalActiviteSerializer(entrees, many=True).data)


def _duree_heures(reservation):
    debut = datetime.combine(reservation.date, reservation.heure_debut)
    fin = datetime.combine(reservation.date, reservation.heure_fin)
    return (fin - debut).total_seconds() / 3600


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def rapports_export_excel(request):
    # Alignée sur la restriction déjà en place côté frontend (route /rapports
    # réservée à l'admin) — à élargir ensemble si tu veux l'ouvrir aux
    # techniciens/chercheurs plus tard.
    if request.user.role != Role.ADMIN:
        return Response({'detail': "Seul un administrateur peut exporter ce rapport."}, status=403)

    date_debut = request.query_params.get('date_debut')
    date_fin = request.query_params.get('date_fin')
    laboratoire_id = request.query_params.get('laboratoire')

    reservations = Reservation.objects.exclude(est_archivee=True)
    if date_debut:
        reservations = reservations.filter(date__gte=date_debut)
    if date_fin:
        reservations = reservations.filter(date__lte=date_fin)
    if laboratoire_id:
        reservations = reservations.filter(laboratoire_id=laboratoire_id)

    reservations = list(reservations.select_related('laboratoire', 'demandeur').prefetch_related('equipements'))

    wb = openpyxl.Workbook()
    entete_font = Font(bold=True, color="FFFFFF")
    entete_fill = PatternFill(start_color="1848D9", end_color="1848D9", fill_type="solid")

    def style_entete(ws):
        for cell in ws[ws.max_row]:
            cell.font = entete_font
            cell.fill = entete_fill

    # --- Résumé ---
    ws = wb.active
    ws.title = "Résumé"
    ws.append(["Rapport d'activité — UMRED Labo"])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f"Période : {date_debut or '—'} au {date_fin or '—'}"])
    if laboratoire_id:
        labo = Laboratoire.objects.filter(id=laboratoire_id).first()
        ws.append([f"Laboratoire : {labo.nom if labo else '—'}"])
    ws.append([])
    ws.append(["Indicateur", "Valeur"])
    style_entete(ws)

    equipements_utilises = {e.id for r in reservations for e in r.equipements.all()}
    ws.append(["Réservations sur la période", len(reservations)])
    ws.append(["Équipements utilisés", len(equipements_utilises)])
    ws.append(["Heures d'utilisation cumulées", round(sum(_duree_heures(r) for r in reservations), 1)])
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 18

    # --- Par laboratoire ---
    ws2 = wb.create_sheet("Par laboratoire")
    ws2.append(["Laboratoire", "Réservations", "Heures cumulées"])
    style_entete(ws2)
    par_labo = {}
    for r in reservations:
        d = par_labo.setdefault(r.laboratoire.nom, {'n': 0, 'h': 0})
        d['n'] += 1
        d['h'] += _duree_heures(r)
    for nom, d in sorted(par_labo.items(), key=lambda x: -x[1]['n']):
        ws2.append([nom, d['n'], round(d['h'], 1)])
    ws2.column_dimensions['A'].width = 34

    # --- Par équipement ---
    ws3 = wb.create_sheet("Par équipement")
    ws3.append(["Équipement", "Laboratoire", "Réservations", "Heures cumulées"])
    style_entete(ws3)
    par_equip = {}
    for r in reservations:
        for e in r.equipements.all():
            d = par_equip.setdefault((e.nom, e.laboratoire.nom), {'n': 0, 'h': 0})
            d['n'] += 1
            d['h'] += _duree_heures(r)
    for (nom, labo_nom), d in sorted(par_equip.items(), key=lambda x: -x[1]['n']):
        ws3.append([nom, labo_nom, d['n'], round(d['h'], 1)])
    ws3.column_dimensions['A'].width = 30
    ws3.column_dimensions['B'].width = 30

    # --- Par type d'utilisateur ---
    ws4 = wb.create_sheet("Par type d'utilisateur")
    ws4.append(["Rôle", "Réservations"])
    style_entete(ws4)
    par_role = {}
    for r in reservations:
        label = ROLE_LABELS.get(r.demandeur.role, r.demandeur.role)
        par_role[label] = par_role.get(label, 0) + 1
    for label, n in sorted(par_role.items(), key=lambda x: -x[1]):
        ws4.append([label, n])
    ws4.column_dimensions['A'].width = 26

    # --- Détail ---
    ws5 = wb.create_sheet("Détail des réservations")
    ws5.append(["Date", "Début", "Fin", "Laboratoire", "Équipement(s)", "Demandeur", "Rôle", "Statut"])
    style_entete(ws5)
    for r in sorted(reservations, key=lambda x: x.date):
        equip = ", ".join(e.nom for e in r.equipements.all()) or "Salle uniquement"
        ws5.append([
            r.date.strftime('%d/%m/%Y'), str(r.heure_debut)[:5], str(r.heure_fin)[:5],
            r.laboratoire.nom, equip, r.demandeur.nom_complet,
            ROLE_LABELS.get(r.demandeur.role, r.demandeur.role), r.get_statut_display(),
        ])
    for col, w in zip('ABCDEFGH', [12, 8, 8, 26, 30, 22, 18, 12]):
        ws5.column_dimensions[col].width = w

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="rapport_umred_labo_{datetime.now():%Y%m%d}.xlsx"'
    wb.save(response)
    return response